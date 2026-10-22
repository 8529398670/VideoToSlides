"""Combined decks: a batch's videos in one deck, in the person's order,
with any of them left out.

Run from the project root:  ./.venv/bin/python -m unittest discover tests
"""
from __future__ import annotations

import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace

from PIL import Image
from pptx import Presentation

from app.core import db
from app.jobs import exports, files
from app.models import combined, jobs
from app.sources import youtube


class CombinedTests(unittest.TestCase):
    def setUp(self) -> None:
        self.tmp = tempfile.TemporaryDirectory()
        root = Path(self.tmp.name)
        self.store = db.open_store(root / "app.db", "0" * 64, False)
        self.cfg = SimpleNamespace(storage_dir=root / "storage", jobs_dir=root / "storage" / "jobs")

    def tearDown(self) -> None:
        self.tmp.cleanup()

    def video(self, title: str, slides: int, size: tuple[int, int] = (1280, 720), user_id: int = 1,
              source_url: str | None = None, playlist: dict | None = None) -> jobs.Job:
        job = jobs.create_job(self.store, user_id, jobs.SOURCE_YOUTUBE, title, batch="Course",
                              source_url=source_url, playlist=playlist)
        jobs.claim_queued(self.store, 1)
        directory = files.slides_dir(self.cfg, job.id)
        directory.mkdir(parents=True)
        records = []
        for n in range(1, slides + 1):
            Image.new("RGB", size, (n * 40, 0, 0)).save(directory / f"{n}.jpg")
            records.append({"n": n, "file": f"{n}.jpg", "start_sec": float(n), "captured_sec": float(n)})
        jobs.finish_job(self.store, job.id, size[0], size[1], 60.0, {}, records)
        return jobs.get_job(self.store, job.id)

    def test_choice_is_remembered_and_updated(self) -> None:
        first = combined.save_choice(self.store, 1, "Course", "Course", [3, 1, 2], [2])
        again = combined.save_choice(self.store, 1, "Course", "Whole course", [1, 2, 3], [])
        self.assertEqual(first.id, again.id)
        self.assertEqual((again.title, again.job_ids, again.left_out), ("Whole course", [1, 2, 3], []))
        self.assertIsNone(combined.find(self.store, 2, "Course"))

    def test_deck_follows_order_and_skips_left_out(self) -> None:
        a = self.video("Lecture 1", 2)
        b = self.video("Lecture 2", 1)
        c = self.video("Lecture 10", 3, size=(1024, 768))
        combo = combined.save_choice(self.store, 1, "Course", "Course", [c.id, b.id, a.id], [b.id])
        deck = exports.build_combined_deck(self.store, self.cfg, combo)

        slides = list(Presentation(str(deck)).slides)
        self.assertEqual(len(slides), 5)
        notes = [slide.notes_slide.notes_text_frame.text.splitlines()[0] for slide in slides]
        self.assertEqual(notes, ["Video 1 of 2: Lecture 10"] + ["Video: Lecture 10"] * 2
                         + ["Video 2 of 2: Lecture 1", "Video: Lecture 1"])
        # The first video (4:3) sets the size; the 16:9 slides are letterboxed.
        picture = slides[3].shapes[0]
        self.assertEqual(picture.width, Presentation(str(deck)).slide_width)
        self.assertGreater(picture.top, 0)

    def test_notes_carry_playlist_and_video_links(self) -> None:
        playlist = youtube.Playlist(title="Course", url="https://www.youtube.com/playlist?list=PL1", list_id="PL1")
        first = youtube.Entry(url="https://www.youtube.com/watch?v=aaa", title="Lecture 1", index=2, video_id="aaa")
        second = youtube.Entry(url="https://www.youtube.com/watch?v=bbb", title="Lecture 2", index=1, video_id="bbb")
        a = self.video("Lecture 1", 2, source_url=first.url, playlist=first.playlist_info(playlist))
        b = self.video("Lecture 2", 1, source_url=second.url, playlist=second.playlist_info(playlist))
        combo = combined.save_choice(self.store, 1, "Course", "Course", [a.id, b.id], [])
        slides = list(Presentation(str(exports.build_combined_deck(self.store, self.cfg, combo))).slides)
        notes = [slide.notes_slide.notes_text_frame.text for slide in slides]
        self.assertTrue(notes[0].startswith(
            "Playlist: Course\nhttps://www.youtube.com/playlist?list=PL1\n\n"
            "Video 1 of 2: Lecture 1\nhttps://www.youtube.com/watch?v=aaa\n"
            "In the playlist (video #2): https://www.youtube.com/watch?v=aaa&list=PL1&index=2\n\nScene Start"))
        self.assertTrue(notes[1].startswith("Video: Lecture 1\nScene Start"))
        self.assertTrue(notes[2].startswith("Video 2 of 2: Lecture 2\nhttps://www.youtube.com/watch?v=bbb\n"
                                            "In the playlist (video #1): https://www.youtube.com/watch?v=bbb&list=PL1&index=1"))
        self.assertNotIn("Playlist:", notes[2])

        single = list(Presentation(str(exports.build_deck(self.cfg, jobs.get_job(self.store, b.id)))).slides)
        self.assertTrue(single[0].notes_slide.notes_text_frame.text.startswith(
            "Playlist: Course\nhttps://www.youtube.com/playlist?list=PL1\n\nVideo: Lecture 2\n"))

    def test_pdf_has_a_bookmark_per_video(self) -> None:
        a = self.video("Lecture 1", 2)
        b = self.video("Lecture 2 – ünïcode", 1)
        combo = combined.save_choice(self.store, 1, "Course", "Whole course", [b.id, a.id], [])
        combined.start_export(self.store, combo.id, [jobs.EXPORT_PDF])
        exports.run_combined_export(self.store, self.cfg, combo.id, [jobs.EXPORT_PDF])
        state = combined.get(self.store, combo.id).exports
        self.assertEqual(state[jobs.EXPORT_PDF]["status"], jobs.EXPORT_DONE)
        self.assertNotIn(jobs.EXPORT_PPTX, state)
        directory = files.combined_dir(self.cfg, combo.id)
        self.assertFalse((directory / files.DECK_FILE).exists())
        data = (directory / files.PDF_FILE).read_bytes()
        self.assertTrue(data.startswith(b"%PDF-1.4"))
        self.assertEqual(data.count(b"/Type /Page "), 3)
        self.assertIn(b"/PageMode /UseOutlines", data)
        self.assertIn(b"/Title (Lecture 1)", data)
        self.assertIn(b"/Title <FEFF" + "Lecture 2 – ünïcode".encode("utf-16-be").hex().upper().encode(), data)

    def test_single_pdf_has_no_bookmarks(self) -> None:
        job = self.video("Solo", 2)
        exports.run_export(self.store, self.cfg, job.id, [jobs.EXPORT_PDF])
        self.assertEqual(jobs.get_job(self.store, job.id).exports[jobs.EXPORT_PDF]["status"], jobs.EXPORT_DONE)
        data = (files.job_dir(self.cfg, job.id) / files.PDF_FILE).read_bytes()
        self.assertEqual(data.count(b"/Type /Page "), 2)
        self.assertNotIn(b"/UseOutlines", data)

    def test_someone_elses_or_unfinished_videos_are_skipped(self) -> None:
        mine = self.video("Mine", 1)
        theirs = self.video("Theirs", 1, user_id=2)
        combo = combined.save_choice(self.store, 1, "Course", "Course", [theirs.id, mine.id, 999], [])
        self.assertEqual([job.id for job in exports.combined_jobs(self.store, combo)], [mine.id])

    def test_export_state_and_deck(self) -> None:
        combo = combined.save_choice(self.store, 1, "Course", "Course", [1], [])
        self.assertIsNotNone(combined.start_export(self.store, combo.id, [jobs.EXPORT_GSLIDES]))
        self.assertIsNone(combined.start_export(self.store, combo.id, [jobs.EXPORT_GSLIDES]))
        self.assertEqual(combined.fail_interrupted(self.store), 1)
        combined.finish_export(self.store, combo.id, jobs.EXPORT_GSLIDES, url="https://deck/c", file_id="c1")
        saved = combined.save_choice(self.store, 1, "Course", "Course", [1, 2], [])
        self.assertEqual(saved.gslides_url, "https://deck/c")
        self.assertIsNone(combined.forget_gslides(self.store, combo.id, "other"))
        self.assertIsNone(combined.forget_gslides(self.store, combo.id, "c1").gslides_url)


if __name__ == "__main__":
    unittest.main()
