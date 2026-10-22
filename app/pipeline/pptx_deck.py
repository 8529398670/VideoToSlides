"""Assemble a .pptx: one full-bleed picture per slide, timing in the notes.

A combined deck (several videos) takes its slide size from the first video;
a slide from a video of another shape is fitted inside it, centred, rather
than stretched, and its notes name the video it came from."""
from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from pptx import Presentation
from pptx.util import Emu

from app.pipeline import tunables as T

EMU_PER_IN = 914400
BLANK_LAYOUT = 6


@dataclass
class DeckSlide:
    image: Path
    start_sec: float
    captured_sec: float
    energy: float
    # The video it came from: its title for the notes, and its frame size
    # when that may differ from the deck's.
    label: str | None = None
    size: tuple[int, int] | None = None
    # Notes that open this slide's notes instead of the label: where the
    # video (and its playlist) can be found, on the first slide of each video.
    intro: str | None = None


# Aspect ratios closer than this are treated as equal and filled edge to edge.
ASPECT_TOLERANCE = 0.01


def _placement(item: DeckSlide, slide_w: int, slide_h: int) -> tuple[int, int, int, int]:
    if not item.size or not item.size[0] or not item.size[1]:
        return 0, 0, slide_w, slide_h
    image_aspect = item.size[0] / item.size[1]
    slide_aspect = slide_w / slide_h
    if abs(image_aspect - slide_aspect) / slide_aspect <= ASPECT_TOLERANCE:
        return 0, 0, slide_w, slide_h
    if image_aspect > slide_aspect:
        height = int(slide_w / image_aspect)
        return 0, (slide_h - height) // 2, slide_w, height
    width = int(slide_h * image_aspect)
    return (slide_w - width) // 2, 0, width, slide_h


def build_pptx(slides: list[DeckSlide], width_px: int, height_px: int, out_pptx: Path) -> int:
    aspect = width_px / height_px
    slide_w_in = T.SLIDE_WIDTH_IN
    slide_h_in = slide_w_in / aspect
    prs = Presentation()
    prs.slide_width = Emu(int(slide_w_in * EMU_PER_IN))
    prs.slide_height = Emu(int(slide_h_in * EMU_PER_IN))
    blank = prs.slide_layouts[BLANK_LAYOUT]

    for item in slides:
        slide = prs.slides.add_slide(blank)
        left, top, width, height = _placement(item, prs.slide_width, prs.slide_height)
        slide.shapes.add_picture(str(item.image), left, top, width=width, height=height)
        notes = slide.notes_slide.notes_text_frame
        timing = (f"Scene Start: {item.start_sec:.2f}s\n"
                  f"Captured at: {item.captured_sec:.2f}s\n"
                  f"Energy: {item.energy:.2f}")
        if item.intro:
            notes.text = f"{item.intro}\n\n{timing}"
        elif item.label:
            notes.text = f"Video: {item.label}\n{timing}"
        else:
            notes.text = timing

    # Written beside the target and renamed, so a download never sees half a file.
    partial = out_pptx.with_suffix(".pptx.part")
    prs.save(str(partial))
    partial.replace(out_pptx)
    return len(slides)
